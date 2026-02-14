#!/usr/bin/env python3
"""
Simplified Job Finder - Single Master File Approach
Scrapes LinkedIn, filters jobs, and appends to jobs_master.xlsx

SAFETY FEATURES:
- Uses absolute paths (works from any directory)
- Creates backup before every write
- Aborts on read errors (never loses data)
- File locking to prevent concurrent writes
"""

import os
import sys
import shutil
import fcntl
import pandas as pd
from datetime import datetime
from pathlib import Path

# Get absolute path to script directory
SCRIPT_DIR = Path(__file__).parent.absolute()

# Add src to path
sys.path.insert(0, str(SCRIPT_DIR / 'src'))

from job_scraper import JobScraper
from simple_evaluator import SimpleEvaluator
from skill_matcher import SkillMatcher

# Configuration - USE ABSOLUTE PATHS
MASTER_FILE = SCRIPT_DIR / "jobs_master.xlsx"
BACKUP_DIR = SCRIPT_DIR / "backups"
LOCK_FILE = SCRIPT_DIR / ".jobs_master.lock"

SEARCH_QUERIES = [
    ("AI Engineer", '("AI Engineer" OR "Artificial Intelligence Engineer") NOT (Senior OR Sr. OR Principal OR Staff OR Lead OR Manager)', 20),
    ("Machine Learning Engineer", '("Machine Learning Engineer" OR "ML Engineer") NOT (Senior OR Sr. OR Principal OR Staff OR Lead OR Manager)', 20),
    ("Data Scientist", '"Data Scientist" AND ("Machine Learning" OR "Deep Learning" OR "AI" OR "Generative AI") NOT (Senior OR Sr. OR Principal OR Staff OR Lead OR Manager)', 20)
]


def acquire_lock():
    """Acquire file lock to prevent concurrent writes."""
    try:
        lock_fd = open(LOCK_FILE, 'w')
        fcntl.flock(lock_fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        return lock_fd
    except (IOError, OSError):
        print("[ERROR] Another instance is running. Wait for it to finish.")
        raise SystemExit(1)


def release_lock(lock_fd):
    """Release file lock."""
    if lock_fd:
        fcntl.flock(lock_fd, fcntl.LOCK_UN)
        lock_fd.close()
        try:
            os.remove(LOCK_FILE)
        except:
            pass


def create_backup():
    """Create a backup of the master file before writing."""
    if not MASTER_FILE.exists():
        return None

    # Create backups directory if needed
    BACKUP_DIR.mkdir(exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = BACKUP_DIR / f"jobs_master_backup_{timestamp}.xlsx"

    try:
        shutil.copy2(MASTER_FILE, backup_path)
        print(f"[BACKUP] Created: {backup_path.name}")

        # Keep only last 10 backups to save space
        backups = sorted(BACKUP_DIR.glob("jobs_master_backup_*.xlsx"))
        if len(backups) > 10:
            for old_backup in backups[:-10]:
                old_backup.unlink()

        return backup_path
    except Exception as e:
        print(f"[ERROR] Failed to create backup: {e}")
        print("[ERROR] Aborting to prevent data loss.")
        raise SystemExit(1)


def load_existing_master():
    """Load existing master file. ABORTS on error to prevent data loss."""
    if not MASTER_FILE.exists():
        print(f"[INFO] No existing file found. Will create new: {MASTER_FILE.name}")
        return pd.DataFrame()

    try:
        # Try "All Jobs" sheet first, then fall back to first sheet
        try:
            df = pd.read_excel(MASTER_FILE, sheet_name="All Jobs")
        except:
            df = pd.read_excel(MASTER_FILE, sheet_name=0)

        if len(df) == 0:
            print(f"[WARNING] Master file exists but has 0 rows!")
            print(f"[WARNING] This might indicate data loss. Check backups in {BACKUP_DIR}")
        else:
            print(f"[INFO] Loaded {len(df)} existing jobs from {MASTER_FILE.name}")

        return df

    except Exception as e:
        print(f"[ERROR] Could not load existing file: {e}")
        print(f"[ERROR] Aborting to prevent data loss.")
        print(f"[ERROR] Check the file manually or restore from backup in {BACKUP_DIR}")
        raise SystemExit(1)


def create_master_excel(df):
    """Create Excel with formatting and dropdown. Creates backup first."""
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    # CRITICAL: Create backup before writing
    create_backup()

    # Validate we have data
    if df.empty:
        print("[ERROR] Refusing to write empty DataFrame!")
        raise SystemExit(1)

    print(f"[INFO] Writing {len(df)} jobs to {MASTER_FILE.name}...")

    # Save to Excel
    with pd.ExcelWriter(MASTER_FILE, engine='openpyxl') as writer:
        # Sheet 1: All Jobs (everything)
        df.to_excel(writer, sheet_name='All Jobs', index=False)

        # Sheet 2: Top Matches (70%+ skill score, sorted by score)
        if 'Keywords_Matching_Score' in df.columns:
            top_matches = df[df['Keywords_Matching_Score'] >= 70].copy()
            top_matches = top_matches.sort_values('Keywords_Matching_Score', ascending=False)
            top_matches.to_excel(writer, sheet_name='Top Matches', index=False)

        # Sheet 3: Applied (only checked applications)
        if 'Applied' in df.columns:
            applied = df[df['Applied'] == 'Applied'].copy()
            applied.to_excel(writer, sheet_name='Applied', index=False)

    # Add formatting
    wb = load_workbook(MASTER_FILE)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Add dropdown to "Applied" column if it exists
        if ws.max_row > 1:
            # Find Applied column
            applied_col = None
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=col).value == 'Applied':
                    applied_col = col
                    break

            if applied_col:
                col_letter = chr(64 + applied_col) if applied_col <= 26 else f"A{chr(64 + applied_col - 26)}"
                dv = DataValidation(type="list", formula1='"Not Applied,Applied"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f'{col_letter}2:{col_letter}{ws.max_row}')

        # Color code rows by Tier
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        orange_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        # Find Tier column
        tier_col = None
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=col).value == 'Tier':
                tier_col = col
                break

        if tier_col:
            for row in range(2, ws.max_row + 1):
                tier_cell = ws.cell(row=row, column=tier_col)
                tier_value = str(tier_cell.value) if tier_cell.value else ""

                if "🟢" in tier_value:
                    fill = green_fill
                elif "🟡" in tier_value:
                    fill = yellow_fill
                elif "🟠" in tier_value:
                    fill = orange_fill
                else:
                    fill = red_fill

                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = fill

        # Set column widths
        column_widths = {
            'A': 6, 'B': 20, 'C': 35, 'D': 15, 'E': 50,
            'F': 15, 'G': 12, 'H': 10, 'I': 18, 'J': 40,
            'K': 40, 'L': 10, 'M': 15, 'N': 20, 'O': 60
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

    wb.save(MASTER_FILE)
    print(f"[SUCCESS] Master file updated: {MASTER_FILE.name}")


def main():
    print("\n" + "="*70)
    print("  🔍 JOB FINDER - Single Master File System")
    print(f"  Working Directory: {SCRIPT_DIR}")
    print("="*70 + "\n")

    # Acquire lock to prevent concurrent runs
    lock_fd = acquire_lock()

    try:
        # Load existing jobs
        existing_df = load_existing_master()
        existing_urls = set(existing_df['Job_Link'].dropna().tolist()) if not existing_df.empty and 'Job_Link' in existing_df.columns else set()

        # Initialize components
        chrome_path = os.getenv("CHROME_PATH", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
        chrome_profile = os.getenv("CHROME_PROFILE_PATH", "")
        blacklist_path = str(SCRIPT_DIR / "config" / "blacklist.txt")

        simple_eval = SimpleEvaluator()
        skill_matcher = SkillMatcher()

        all_new_jobs = []
        scraper = None

        # Run searches
        for role_name, keywords, count in SEARCH_QUERIES:
            print(f"\n[SEARCH] {role_name} (Target: {count} jobs)")
            print("-" * 70)

            # Create new scraper for each search to avoid browser reuse issues
            scraper = JobScraper(chrome_path, chrome_profile, blacklist_path, max_jobs=count)
            jobs = scraper.search_and_collect(keywords)
            scraper.close_browser()

            if not jobs:
                print(f"[INFO] No jobs found for {role_name}")
                continue

            # Filter and evaluate
            for job in jobs:
                # Skip if already in master file
                if job.job_link in existing_urls:
                    continue

                # Simple filter (Senior/Citizenship)
                passed, reason = simple_eval.quick_check(job.job_title, job.job_description)

                if not passed:
                    verdict = "NO"
                    verdict_reason = f"Rejected: {reason}"
                    skill_score = 0
                    match_count = "0/0"
                    matched_skills = ""
                    missing_skills = ""
                    tier = "🔴 Filtered Out"
                else:
                    score_result = skill_matcher.score_job(
                        job.job_description,
                        is_premium=job.is_premium,
                        is_entry_level=(job.role_type == "Entry-level")
                    )

                    skill_score = score_result['final_score']
                    match_count = score_result['match_count']
                    matched_skills = ", ".join(score_result['matched_skills'][:10])
                    missing_skills = ", ".join(score_result['missing_skills'][:10])
                    tier = score_result['tier']

                    verdict = "YES"
                    verdict_reason = f"Skill Match: {skill_score}% ({match_count})"

                all_new_jobs.append({
                    'Company': job.company_name,
                    'Job_Title': job.job_title,
                    'Location': job.location,
                    'Job_Link': job.job_link,
                    'Date_Added': datetime.now().strftime("%Y-%m-%d %H:%M"),
                    'Keywords_Matching_Score': skill_score,
                    'Match': match_count,
                    'Tier': tier,
                    'Matched Skills': matched_skills,
                    'Missing Skills': missing_skills,
                    'Verdict': verdict,
                    'Reason': verdict_reason,
                    'Applied': 'Not Applied',
                    'Search_Query': job.search_query,
                    'Job_Description': job.job_description[:2000] if job.job_description else ''
                })

        # Merge with existing
        if all_new_jobs:
            new_df = pd.DataFrame(all_new_jobs)

            # CRITICAL: Only combine if we have existing data or new data
            if not existing_df.empty:
                print(f"[INFO] Combining {len(existing_df)} existing + {len(new_df)} new jobs")
                combined_df = pd.concat([existing_df, new_df], ignore_index=True)
            else:
                combined_df = new_df

            # Create master file (with backup)
            create_master_excel(combined_df)

            # Calculate statistics
            total_new = len(all_new_jobs)
            passed_filter = len([j for j in all_new_jobs if j['Verdict'] == 'YES'])
            good_matches = len([j for j in all_new_jobs if j['Keywords_Matching_Score'] >= 70])
            perfect_matches = len([j for j in all_new_jobs if j['Keywords_Matching_Score'] >= 90])

            print(f"\n[SUCCESS] Added {total_new} new jobs to {MASTER_FILE.name}")
            print(f"[INFO] Total jobs in master: {len(combined_df)}")
            print(f"\n📊 STATISTICS:")
            print(f"  • Total fetched: {total_new}")
            print(f"  • Passed filters (YES): {passed_filter} ({int(passed_filter/total_new*100) if total_new > 0 else 0}%)")
            print(f"  • Good matches (70%+): {good_matches} ({int(good_matches/total_new*100) if total_new > 0 else 0}%)")
            print(f"  • Perfect matches (90%+): {perfect_matches} ({int(perfect_matches/total_new*100) if total_new > 0 else 0}%)")
        else:
            print("\n[INFO] No new jobs found (all were duplicates)")

        print("\n" + "="*70)
        print("  ✅ COMPLETE - Open jobs_master.xlsx to review")
        print("="*70 + "\n")

    finally:
        # Always release lock
        release_lock(lock_fd)
        if scraper:
            try:
                scraper.close_browser()
            except:
                pass


if __name__ == "__main__":
    main()
